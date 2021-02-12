from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart
wb = load_workbook("Sample.xlsx")
ws = wb.active

###### BarChart ######
# Reference : 어떤 data를 차트로 만들지에 대한 정의
# 영어, 수학 성적을 차트로 그려보자
# B2:C11 까지의 데이터를 차트로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()  # 차트 종류 설정
# bar_chart.add_data(bar_value)  # 차트 데이터 추가
# ws.add_chart(bar_chart, "E1")  # 위에서 다 그려진 차트를 worksheet안에 넣어줌
#######################

###### LineChart ######
# B1:C11 까지의 데이터를 차트로 생성
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True)  # 계열 이름을 data값 그대로 가져옴
line_chart.title = "성적표"  # 제목
line_chart.style = 20  # 미리 정의된 스타일 적용
line_chart.y_axis.title = "점수"  # Y축 제목 설정
line_chart.x_axis.title = "번호"  # X축 제목 설정
ws.add_chart(line_chart, "E1")
#######################

wb.save("Sample_chart.xlsx")
